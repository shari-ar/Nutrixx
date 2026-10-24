"""Extract identified Ciqual 2025 observations from official English and XML files."""

import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook


NUMERIC = re.compile(r"^(?:0|[1-9]\d*)(?:[.,]\d+)?$")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cells(element):
    return {child.tag: (child.text or "").strip() for child in element}


def xml_rows(path, row_tag):
    for _, element in ET.iterparse(path, events=("end",)):
        if element.tag == row_tag:
            yield cells(element)
            element.clear()


def main():
    if len(sys.argv) != 6:
        raise SystemExit(
            "Usage: extract-ciqual-2025.py <english.xlsx> <compo.xml> "
            "<const.xml> <sources.xml> <output.json>"
        )
    workbook_path, composition_path, components_path, sources_path, output_path = map(
        Path, sys.argv[1:]
    )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["food composition"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        if len(header) != 84:
            raise ValueError("Ciqual English workbook column count changed")
        foods = {}
        for row in rows:
            code = str(row[6]).strip()
            name = row[7]
            if not code.isdecimal() or not isinstance(name, str) or code in foods:
                raise ValueError(f"Invalid or repeated Ciqual food identity: {code}")
            foods[code] = {
                "code": code,
                "name": name.strip(),
                "group": str(row[3] or "").strip(),
                "values": {},
            }
        if len(foods) != 3484:
            raise ValueError(f"Unexpected Ciqual food count: {len(foods)}")
    finally:
        workbook.close()

    components = {}
    for row in xml_rows(components_path, "CONST"):
        code = row["const_code"]
        if code in components:
            raise ValueError(f"Repeated Ciqual component code: {code}")
        components[code] = {
            "name": row["const_nom_eng"],
            "infoods": row["code_INFOODS"],
        }

    sources = {}
    for row in xml_rows(sources_path, "SOURCES"):
        code = row["source_code"]
        if code in sources:
            raise ValueError(f"Repeated Ciqual evidence source: {code}")
        sources[code] = row["ref_citation"]

    counts = {"xmlRows": 0, "numeric": 0, "missing": 0, "traces": 0, "bounded": 0}
    for row in xml_rows(composition_path, "COMPO"):
        counts["xmlRows"] += 1
        food = foods.get(row["alim_code"])
        if food is None:
            raise ValueError(f"Unknown Ciqual food code: {row['alim_code']}")
        component_code = row["const_code"]
        if component_code not in components or component_code in food["values"]:
            raise ValueError(f"Unknown or repeated Ciqual component: {component_code}")
        raw = row["teneur"]
        if raw == "-":
            counts["missing"] += 1
        elif raw.lower() == "traces":
            counts["traces"] += 1
        elif raw.startswith("<"):
            counts["bounded"] += 1
        elif NUMERIC.fullmatch(raw):
            counts["numeric"] += 1
        else:
            raise ValueError(f"Unexpected Ciqual value for {component_code}: {raw}")
        confidence = row["code_confiance"]
        if confidence not in ("", "A", "B", "C", "D"):
            raise ValueError(f"Unexpected Ciqual confidence: {confidence}")
        source_code = row["source_code"]
        if source_code and source_code not in sources:
            raise ValueError(f"Unknown Ciqual evidence source: {source_code}")
        food["values"][component_code] = {
            "raw": raw,
            "confidence": confidence,
            "sourceCode": source_code,
        }

    output = {
        "format": "nutrixx.ciqual-source-extract",
        "release": "2025",
        "sourceSha256": {
            "workbook": digest(workbook_path),
            "composition": digest(composition_path),
            "components": digest(components_path),
            "sources": digest(sources_path),
        },
        "counts": counts,
        "components": components,
        "sources": sources,
        "foods": list(foods.values()),
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(output, ensure_ascii=False), encoding="utf-8")
    print(f"Extracted {len(foods)} Ciqual foods and {counts['xmlRows']} observations")


if __name__ == "__main__":
    main()
