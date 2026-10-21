"""Extract AFCD Release 3 workbook rows without changing the source files."""

import json
import hashlib
import sys
from pathlib import Path

from openpyxl import load_workbook


def records(path: Path, sheet_name: str):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        next(rows)
        headers = next(rows)
        if headers[0] != "Public Food Key":
            raise ValueError(f"Unexpected AFCD header in {path.name}")
        for index, row in enumerate(rows, start=4):
            if isinstance(row[0], str) and row[0].startswith("F"):
                yield index, headers, row
    finally:
        workbook.close()


def main():
    if len(sys.argv) != 4:
        raise SystemExit(
            "Usage: extract-afcd-release.py <food-details.xlsx> "
            "<nutrient-profiles.xlsx> <output.json>"
        )
    details_path, profiles_path, output_path = map(Path, sys.argv[1:])
    foods = {}
    for row_number, _, row in records(details_path, "Food details"):
        key = row[0]
        if key in foods:
            raise ValueError(f"Duplicate AFCD food key: {key}")
        foods[key] = {
            "key": key,
            "detailRow": row_number,
            "derivation": row[2],
            "name": row[3],
            "samplingDetails": row[5],
        }

    profiles = {}
    for row_number, headers, row in records(
        profiles_path, "All solids & liquids per 100 g"
    ):
        key = row[0]
        if key in profiles:
            raise ValueError(f"Duplicate AFCD profile key: {key}")
        values = {}
        for column, (header, value) in enumerate(zip(headers, row), start=1):
            if column <= 4 or not isinstance(value, (int, float)):
                continue
            if not isinstance(header, str):
                raise ValueError(f"Missing AFCD column {column} header")
            values[str(column)] = value
        profiles[key] = {"profileRow": row_number, "values": values}

    if set(foods) != set(profiles):
        raise ValueError("AFCD food details and nutrient profiles have different keys")
    payload = {
        "format": "nutrixx.afcd-source-extract",
        "release": "3",
        "sourceSha256": {
            "foodDetails": hashlib.sha256(details_path.read_bytes()).hexdigest(),
            "nutrientProfiles": hashlib.sha256(profiles_path.read_bytes()).hexdigest(),
        },
        "columns": {
            str(index): header
            for index, header in enumerate(headers, start=1)
            if index > 4 and isinstance(header, str)
        },
        "foods": [{**foods[key], **profiles[key]} for key in sorted(foods)],
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print(f"Extracted {len(foods)} AFCD foods to {output_path}")


if __name__ == "__main__":
    main()
