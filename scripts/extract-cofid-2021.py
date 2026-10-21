"""Extract aligned CoFID 2021 records without changing the official workbook."""

import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


SHEETS = (
    "1.3 Proximates",
    "1.4 Inorganics",
    "1.5 Vitamins",
    "1.6 Vitamin Fractions",
    "1.8 (SFA per 100gFood)",
    "1.10 (MUFA per 100gFood)",
    "1.12 (PUFA per 100gFood)",
    "1.13 Phytosterols",
)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: extract-cofid-2021.py <CoFID-2021.xlsx> <output.json>")
    source_path, output_path = map(Path, sys.argv[1:])
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        primary = workbook[SHEETS[0]]
        rows = list(primary.iter_rows(min_row=4, values_only=True))
        if len(rows) != 2887:
            raise ValueError(f"Unexpected CoFID food row count: {len(rows)}")
        codes = Counter(str(row[0]).strip() for row in rows)
        foods = []
        for number, row in enumerate(rows, start=4):
            code, name, description, group = row[:4]
            if not isinstance(code, str) or not isinstance(name, str):
                raise ValueError(f"Invalid CoFID food identity at row {number}")
            code = code.strip()
            foods.append(
                {
                    "key": f"{code}@{number}" if codes[code] > 1 else code,
                    "code": code,
                    "row": number,
                    "name": name.strip(),
                    "description": description,
                    "group": group,
                    "mainReference": row[5],
                    "values": {},
                }
            )

        columns = {}
        for sheet_name in SHEETS:
            sheet = workbook[sheet_name]
            records = sheet.iter_rows(values_only=True)
            headings = next(records)
            codes_row = next(records)
            next(records)  # Descriptions; the primary header retains the source label.
            if sheet.max_row != 2890:
                raise ValueError(f"Unexpected CoFID sheet length: {sheet_name}")
            sheet_columns = {}
            for index, code in enumerate(codes_row):
                if index < 7 or not isinstance(code, str) or not code.strip():
                    continue
                key = f"{sheet_name}:{code.strip()}"
                if key in columns:
                    raise ValueError(f"Duplicate CoFID column: {key}")
                columns[key] = str(headings[index]).strip()
                sheet_columns[index] = key
            for index, row in enumerate(records):
                food = foods[index]
                if (str(row[0]).strip(), row[1]) != (food["code"], food["name"]):
                    raise ValueError(f"CoFID sheet alignment changed: {sheet_name}, row {index + 4}")
                for position, key in sheet_columns.items():
                    value = row[position]
                    if value is not None and value != "":
                        food["values"][key] = value

        output = {
            "format": "nutrixx.cofid-source-extract",
            "release": "2021",
            "sourceSha256": hashlib.sha256(source_path.read_bytes()).hexdigest(),
            "duplicateFoodCodes": sorted(code for code, count in codes.items() if count > 1),
            "columns": columns,
            "foods": foods,
        }
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(output, ensure_ascii=False), encoding="utf-8")
        print(f"Extracted {len(foods)} CoFID rows to {output_path}")
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
